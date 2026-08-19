import openpyxl
import json
import os
from datetime import datetime

ARQUIVOS = [
    "/home/ubuntu/upload/RelatóriodeControledeFériasdeFuncionáriosSOLAR(1).xlsx",
    "/home/ubuntu/upload/RelatóriodeControledeFériasdeFuncionáriosFREIRE(1).xlsx",
    "/home/ubuntu/upload/RelatóriodeControledeFériasdeFuncionáriosJOANES(1).xlsx",
    "/home/ubuntu/upload/RelatóriodeControledeFériasdeFuncionáriosSUDOESTE(1).xlsx",
]

def fmt_date(val):
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d")
    if isinstance(val, str):
        val = val.strip()
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
            try:
                return datetime.strptime(val, fmt).strftime("%Y-%m-%d")
            except:
                pass
    return str(val) if val else None

todos = []

for arquivo in ARQUIVOS:
    empresa = os.path.basename(arquivo).replace("RelatóriodeControledeFériasdeFuncionários", "").replace("(1).xlsx", "").strip()
    wb = openpyxl.load_workbook(arquivo, data_only=True)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        # Encontrar a linha de cabeçalho
        header_row = None
        header_idx = {}
        for i, row in enumerate(ws.iter_rows(values_only=True), 1):
            row_str = [str(c).strip().upper() if c else "" for c in row]
            # Procurar colunas relevantes
            for j, cell in enumerate(row_str):
                if "NOME" in cell and "COLABORADOR" in cell:
                    header_row = i
                    header_idx["nome"] = j
                elif "NOME" in cell and header_row is None:
                    header_row = i
                    header_idx["nome"] = j
            if header_row:
                for j, cell in enumerate(row_str):
                    if "ADMISS" in cell:
                        header_idx["admissao"] = j
                    elif "PERÍODO" in cell or "PERIODO" in cell or "AQUISITIVO" in cell:
                        header_idx["periodoAquisitivo"] = j
                    elif "VENCIMENTO" in cell:
                        header_idx["vencimento"] = j
                    elif "LIMITE" in cell or "DATA LIMITE" in cell:
                        header_idx["dataLimite"] = j
                    elif "SALDO" in cell:
                        header_idx["saldo"] = j
                    elif "DIREITO" in cell or "DIAS DIREITO" in cell:
                        header_idx["diasDireito"] = j
                break
        
        if not header_row or "nome" not in header_idx:
            print(f"[AVISO] Cabeçalho não encontrado em {empresa} / {sheet_name}")
            continue
        
        print(f"\n=== {empresa} / {sheet_name} ===")
        print(f"Colunas encontradas: {header_idx}")
        
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            nome_val = row[header_idx["nome"]] if header_idx.get("nome") is not None and len(row) > header_idx["nome"] else None
            if not nome_val or str(nome_val).strip() == "" or str(nome_val).strip().upper() in ("NOME", "NOME DO COLABORADOR"):
                continue
            
            nome = str(nome_val).strip().upper()
            
            rec = {
                "empresa": empresa,
                "nome": nome,
                "admissao": fmt_date(row[header_idx["admissao"]]) if "admissao" in header_idx and len(row) > header_idx["admissao"] else None,
                "periodoAquisitivo": fmt_date(row[header_idx["periodoAquisitivo"]]) if "periodoAquisitivo" in header_idx and len(row) > header_idx["periodoAquisitivo"] else None,
                "vencimento": fmt_date(row[header_idx["vencimento"]]) if "vencimento" in header_idx and len(row) > header_idx["vencimento"] else None,
                "dataLimite": fmt_date(row[header_idx["dataLimite"]]) if "dataLimite" in header_idx and len(row) > header_idx["dataLimite"] else None,
                "saldo": row[header_idx["saldo"]] if "saldo" in header_idx and len(row) > header_idx["saldo"] else None,
                "diasDireito": row[header_idx["diasDireito"]] if "diasDireito" in header_idx and len(row) > header_idx["diasDireito"] else None,
            }
            todos.append(rec)
            print(f"  {nome} | adm:{rec['admissao']} | pa:{rec['periodoAquisitivo']} | venc:{rec['vencimento']} | lim:{rec['dataLimite']}")

with open("/home/ubuntu/ferias-manager/scripts/dados_planilhas.json", "w", encoding="utf-8") as f:
    json.dump(todos, f, ensure_ascii=False, indent=2)

print(f"\nTotal de registros extraídos: {len(todos)}")
